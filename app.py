import streamlit as st
import pandas as pd
import time
import io
import os
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import plotly.express as px
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv


load_dotenv()


st.set_page_config(
    page_title="Consolidador Financeiro & Vendas",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Consolidador de Relatórios Financeiros e Vendas")
st.markdown("""
Esta ferramenta realiza o processo de **ETL (Extração, Tratamento e Carga)** automatizado.
Faça o upload das planilhas de quantos setores quiser para gerar o relatório consolidado de forma instantânea.
""")

st.divider()

st.sidebar.header("📁 Upload das Planilhas")

uploaded_files = st.sidebar.file_uploader(
    "Selecione ou arraste as planilhas (.csv ou .xlsx)",
    type=["csv", "xlsx"],
    accept_multiple_files=True
)

st.sidebar.divider()
st.sidebar.header("📧 Envio por E-mail")
email_destinatario = st.sidebar.text_input("E-mail do Destinatário", placeholder="diretoria@empresa.com")



def enviar_email(destinatario, arquivo_bytes, nome_arquivo):
    SMTP_SERVER = "smtp.gmail.com"
    SMTP_PORT = 465
    
    
    REMETENTE_EMAIL = os.getenv("GMAIL_USER")
    REMETENTE_SENHA = os.getenv("GMAIL_APP_PASSWORD")

    if not REMETENTE_EMAIL or not REMETENTE_SENHA:
        st.error("Erro: Credenciais de e-mail não encontradas no arquivo .env!")
        return False

    try:
        msg = MIMEMultipart()
        msg['From'] = REMETENTE_EMAIL
        msg['To'] = destinatario
        msg['Subject'] = "📊 Relatório Consolidado de Vendas e Finanças Automatizado"

        corpo = """
        Olá,

        Segue em anexo o relatório consolidado contendo os dados higienizados e atualizados dos diferentes setores.

        Este e-mail foi gerado e enviado automaticamente via Pipeline de Automação Python.

        Atenciosamente,
        Equipe de Automação & Dados
        """
        msg.attach(MIMEText(corpo, 'plain'))

        part = MIMEBase('application', 'octet-stream')
        part.set_payload(arquivo_bytes)
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', f'attachment; filename="{nome_arquivo}"')
        msg.attach(part)

        server = smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT, timeout=15)
        server.login(REMETENTE_EMAIL, REMETENTE_SENHA)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        st.error(f"Erro ao enviar e-mail: {e}")
        return False



def processar_e_consolidar(lista_arquivos):
    dfs = []
    
    for file in lista_arquivos:
        if file.name.endswith('.csv'):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
        
        df.columns = df.columns.str.strip()
        
        colunas_renomear = {}
        for col in df.columns:
            nome_lower = col.lower().replace(' ', '_').replace('-', '_')
            if any(k in nome_lower for k in ['data', 'dt_venda', 'dt_emissao', 'date']):
                colunas_renomear[col] = 'Data_Venda'
        
        if colunas_renomear:
            df.rename(columns=colunas_renomear, inplace=True)

        nome_setor = file.name.split('.')[0].replace('_', ' ').title()
        df['Origem / Setor'] = nome_setor
        dfs.append(df)

    df_consolidado = pd.concat(dfs, ignore_index=True)
    df_consolidado.drop_duplicates(inplace=True)
    
    if 'Data_Venda' in df_consolidado.columns:
        df_consolidado['Data_Venda'] = pd.to_datetime(df_consolidado['Data_Venda'], errors='coerce')
        df_consolidado['Data_Venda'] = df_consolidado['Data_Venda'].ffill().dt.date

    for col in df_consolidado.columns:
        if any(k in col.lower() for k in ['valor', 'total', 'receita', 'vendas', 'desconto']):
            df_consolidado[col] = pd.to_numeric(df_consolidado[col], errors='coerce').fillna(0)

    return df_consolidado


if 'dados_processados' not in st.session_state:
    st.session_state.dados_processados = False
    st.session_state.df_final = None
    st.session_state.tempo_execucao = 0
    st.session_state.excel_data = None



if uploaded_files:
    st.sidebar.success(f"📎 {len(uploaded_files)} arquivo(s) carregado(s)!")
    
    if st.button("🚀 Processar e Consolidar Relatórios", type="primary"):
        inicio = time.time()
        
        with st.spinner("Higienizando dados e consolidando relatórios..."):
            df_final = processar_e_consolidar(uploaded_files)
            tempo_execucao = time.time() - inicio

           
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                df_final.to_excel(writer, index=False, sheet_name='Consolidado')
                worksheet = writer.sheets['Consolidado']
                
                header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                
                for col_num, col_name in enumerate(df_final.columns, 1):
                    cell = worksheet.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    if any(k in col_name.lower() for k in ['valor', 'total', 'receita', 'desconto']):
                        for row in range(2, len(df_final) + 2):
                            worksheet.cell(row=row, column=col_num).number_format = 'R$ #,##0.00'
                    elif 'data' in col_name.lower():
                        for row in range(2, len(df_final) + 2):
                            cell_data = worksheet.cell(row=row, column=col_num)
                            cell_data.number_format = 'DD/MM/YYYY'
                            cell_data.alignment = Alignment(horizontal="center")
                            
                    col_letter = get_column_letter(col_num)
                    max_data_len = df_final[col_name].fillna("").astype(str).str.len().max()
                    max_len = max(max_data_len if pd.notna(max_data_len) else 0, len(col_name)) + 4
                    worksheet.column_dimensions[col_letter].width = int(max_len)

            st.session_state.df_final = df_final
            st.session_state.tempo_execucao = tempo_execucao
            st.session_state.excel_data = buffer.getvalue()
            st.session_state.dados_processados = True

    if st.session_state.dados_processados:
        df_final = st.session_state.df_final
        tempo_execucao = st.session_state.tempo_execucao
        excel_data = st.session_state.excel_data
        nome_arquivo = "Relatorio_Consolidado_Vendas_Financas.xlsx"

        st.success("✅ Relatório consolidado com sucesso!")

       
        st.subheader("⚡ Impacto e Ganho de Eficiência")
        col_tempo1, col_tempo2, col_tempo3 = st.columns(3)
        
        horas_estimadas = len(uploaded_files) * 1
        col_tempo1.metric(label="Tempo Estimado Manual", value=f"~{horas_estimadas} Horas")
        col_tempo2.metric(label="Tempo com Automação Python", value=f"{tempo_execucao:.2f} Segundos")
        col_tempo3.metric(label="Ganho de Eficiência", value="99.9%", delta="Automação Completa")

        st.divider()

        # --- DASHBOARD EXECUTIVO ---
        st.subheader("📈 Dashboard Executivo Consolidado")
        
        col_kpi1, col_kpi2, col_kpi3 = st.columns(3)
        col_kpi1.metric("Total de Registros", f"{len(df_final):,} linhas")
        
        colunas_numericas = df_final.select_dtypes(include=['float64', 'int64']).columns
        if len(colunas_numericas) > 0:
            col_kpi2.metric("Soma Total", f"R$ {df_final[colunas_numericas[0]].sum():,.2f}")
            col_kpi3.metric("Média por Registro", f"R$ {df_final[colunas_numericas[0]].mean():,.2f}")

            st.markdown("#### Distribuição por Origem/Setor")
            fig = px.bar(
                df_final, 
                x='Origem / Setor', 
                y=colunas_numericas[0], 
                color='Origem / Setor', 
                title=f"Total de {colunas_numericas[0]} por Setor",
                template="plotly_white"
            )
            st.plotly_chart(fig, use_container_width=True)

        st.markdown("#### Prévia dos Dados Tratados")
        st.dataframe(df_final.head(50), use_container_width=True)

    
        col_btn1, col_btn2 = st.columns(2)
        
        with col_btn1:
            st.download_button(
                label="📥 Baixar Relatório Consolidado (.xlsx)",
                data=excel_data,
                file_name=nome_arquivo,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True
            )

        with col_btn2:
            if email_destinatario:
                if st.button("✉️ Enviar Relatório por E-mail", use_container_width=True):
                    with st.spinner("Enviando e-mail..."):
                        sucesso = enviar_email(email_destinatario, excel_data, nome_arquivo)
                        if sucesso:
                            st.success(f"E-mail enviado com sucesso para **{email_destinatario}**!")
            else:
                st.info("💡 Insira um e-mail no menu lateral para habilitar o disparo automático.")

else:
    st.session_state.dados_processados = False
    st.info("👈 Por favor, faça o upload de 1 ou mais planilhas no menu lateral para iniciar o processamento.")